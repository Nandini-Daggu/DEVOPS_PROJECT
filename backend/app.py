"""
AI-Powered Quiz Generator - Backend API
Parses PowerPoint files and generates MCQs using OpenAI
"""

import os
import json
import re
from flask import Flask, request, jsonify
from flask_cors import CORS
from pptx import Presentation
from openai import OpenAI
from dotenv import load_dotenv
import tempfile
import uuid

load_dotenv()

app = Flask(__name__)
CORS(app)

# Configure upload folder
UPLOAD_FOLDER = tempfile.gettempdir()
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max

# Initialize OpenAI client
client = OpenAI(api_key=os.getenv('OPENAI_API_KEY'))

def extract_text_from_pptx(file_path):
    """Extract all text content from a PowerPoint file."""
    try:
        prs = Presentation(file_path)
        slides_text = []
        slide_count = len(prs.slides)
        
        for i, slide in enumerate(prs.slides, 1):
            slide_content = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells]
                        slide_content.append(" | ".join(row_text))
            
            if slide_content:
                slides_text.append({
                    "slide_number": i,
                    "content": "\n".join(slide_content)
                })
        
        return slides_text, slide_count
    except Exception as e:
        raise Exception(f"Error parsing PPTX: {str(e)}")

def generate_mcq_with_ai(slide_content, num_questions, difficulty):
    """Generate MCQs using OpenAI based on slide content."""
    
    difficulty_prompts = {
        "simple": "Generate SIMPLE recall-based questions that test basic understanding of key terms and concepts.",
        "medium": "Generate MEDIUM difficulty questions that test comprehension and application of concepts.",
        "complex": "Generate COMPLEX questions that require analysis, synthesis, and evaluation of the material."
    }
    
    prompt = f"""You are an expert quiz generator. Based on the following PowerPoint slide content, generate {num_questions} multiple-choice questions.

Difficulty Level: {difficulty_prompts.get(difficulty, difficulty_prompts["medium"])}

Slide Content:
{slide_content}

IMPORTANT RULES:
1. Each question must have exactly 4 options labeled A, B, C, D
2. Only ONE correct answer per question
3. Distractors must be plausible but clearly incorrect
4. Questions must be directly relevant to the slide content
5. Include a brief explanation for why the correct answer is right

Return the response in this EXACT JSON format (no markdown, no code blocks):
{{
    "questions": [
        {{
            "id": 1,
            "question": "Question text here?",
            "options": {{
                "A": "Option A text",
                "B": "Option B text",
                "C": "Option C text",
                "D": "Option D text"
            }},
            "correct_answer": "A",
            "explanation": "Explanation of why A is correct and others are wrong"
        }}
    ]
}}"""

    try:
        response = client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "You are a quiz generation expert. Always respond with valid JSON only, no markdown formatting."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.7,
            max_tokens=4000
        )
        
        content = response.choices[0].message.content.strip()
        
        # Clean the response - remove markdown code blocks if present
        content = re.sub(r'```json\s*', '', content)
        content = re.sub(r'```\s*', '', content)
        content = content.strip()
        
        # Parse JSON
        quiz_data = json.loads(content)
        
        # Validate structure
        if "questions" not in quiz_data:
            raise ValueError("Response missing 'questions' key")
        
        for q in quiz_data["questions"]:
            if not all(k in q for k in ["id", "question", "options", "correct_answer", "explanation"]):
                raise ValueError(f"Question {q.get('id', 'unknown')} missing required fields")
            if len(q["options"]) != 4:
                raise ValueError(f"Question {q['id']} does not have exactly 4 options")
        
        return quiz_data["questions"]
    
    except json.JSONDecodeError as e:
        raise Exception(f"Failed to parse AI response as JSON: {str(e)}. Raw response: {content[:500]}")
    except Exception as e:
        raise Exception(f"AI generation error: {str(e)}")

@app.route('/api/health', methods=['GET'])
def health_check():
    """Health check endpoint."""
    return jsonify({"status": "healthy", "message": "Quiz Generator API is running"})

@app.route('/api/upload', methods=['POST'])
def upload_pptx():
    """Upload and parse a PowerPoint file."""
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    
    if not file.filename.lower().endswith(('.ppt', '.pptx')):
        return jsonify({"error": "Only .ppt and .pptx files are accepted"}), 400
    
    try:
        # Save file temporarily
        file_id = str(uuid.uuid4())
        temp_path = os.path.join(app.config['UPLOAD_FOLDER'], f"{file_id}_{file.filename}")
        file.save(temp_path)
        
        # Extract text
        slides_text, slide_count = extract_text_from_pptx(temp_path)
        
        # Clean up temp file
        os.remove(temp_path)
        
        if not slides_text:
            return jsonify({"error": "No text content found in the PowerPoint file"}), 400
        
        return jsonify({
            "file_id": file_id,
            "filename": file.filename,
            "slide_count": slide_count,
            "slides": slides_text
        })
    
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/generate-quiz', methods=['POST'])
def generate_quiz():
    """Generate MCQs from slide content."""
    data = request.json
    
    if not data:
        return jsonify({"error": "No data provided"}), 400
    
    slides = data.get('slides', [])
    num_questions = data.get('num_questions', 10)
    difficulty = data.get('difficulty', 'medium')
    
    if not slides:
        return jsonify({"error": "No slide content provided"}), 400
    
    if num_questions < 5 or num_questions > 30:
        return jsonify({"error": "Number of questions must be between 5 and 30"}), 400
    
    if difficulty not in ['simple', 'medium', 'complex']:
        return jsonify({"error": "Difficulty must be 'simple', 'medium', or 'complex'"}), 400
    
    try:
        # Combine all slide content
        combined_content = "\n\n".join([
            f"--- Slide {s['slide_number']} ---\n{s['content']}" 
            for s in slides
        ])
        
        # Generate questions
        questions = generate_mcq_with_ai(combined_content, num_questions, difficulty)
        
        return jsonify({
            "success": True,
            "questions": questions,
            "total_questions": len(questions),
            "difficulty": difficulty
        })
    
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/validate-quiz', methods=['POST'])
def validate_quiz():
    """Validate quiz answers and return scores with explanations."""
    data = request.json
    
    if not data:
        return jsonify({"error": "No data provided"}), 400
    
    questions = data.get('questions', [])
    answers = data.get('answers', {})
    
    if not questions or not answers:
        return jsonify({"error": "Questions and answers are required"}), 400
    
    results = []
    score = 0
    
    for q in questions:
        q_id = str(q['id'])
        user_answer = answers.get(q_id, None)
        is_correct = user_answer == q['correct_answer']
        
        if is_correct:
            score += 1
        
        results.append({
            "id": q['id'],
            "question": q['question'],
            "options": q['options'],
            "correct_answer": q['correct_answer'],
            "user_answer": user_answer,
            "is_correct": is_correct,
            "explanation": q['explanation']
        })
    
    percentage = round((score / len(questions)) * 100, 1)
    
    return jsonify({
        "score": score,
        "total": len(questions),
        "percentage": percentage,
        "results": results,
        "passed": percentage >= 40
    })

if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=True)